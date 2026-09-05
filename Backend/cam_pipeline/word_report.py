from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import re
from typing import Any, Optional

from cam_pipeline.company_selector import (
    DEFAULT_ARTICLE_CHAR_LIMIT,
    DEFAULT_PROVIDER,
)
from cam_pipeline.chart_analysis import add_chart_analysis_to_plots
from cam_pipeline.market_data import fetch_market_data_for_companies
from cam_pipeline.price_plot import (
    DEFAULT_DAILY_PLOT_INTERVAL,
    DEFAULT_DAILY_PLOT_PERIOD,
    DEFAULT_DAILY_RECENT_POINTS,
    DEFAULT_PLOT_INTERVAL,
    DEFAULT_PLOT_PERIOD,
    DEFAULT_RECENT_POINTS,
    generate_report_price_plots,
)
from cam_pipeline.research_report import (
    DEFAULT_INTERVAL,
    DEFAULT_PERIOD,
    DEFAULT_RECENT_ROWS,
    fetch_selection_research_report,
)

try:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    Document = None
    WD_ALIGN_VERTICAL = None
    WD_TABLE_ALIGNMENT = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    qn = None
    Inches = None
    Pt = None
    RGBColor = None


DEFAULT_WORD_REPORT_OUTPUT_DIR = "word_reports"
DEFAULT_WORD_REPORT_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / DEFAULT_WORD_REPORT_OUTPUT_DIR / "Style.docx"
)
DEFAULT_BODY_FONT = "NanumSquareOTF"
MIRAE_ORANGE = "F58220"
MIRAE_BLACK = "111111"
MIRAE_LIGHT_BLACK = "222222"
MIRAE_DIVIDER = "8A8A8A"
MIRAE_BEIGE = "E7E0D6"
MIRAE_PAPER = "F6F2EC"
MIRAE_SOFT_GRAY = "6E6E6E"
HAF_GREEN = "1E7032"
HAF_DEEP_GREEN = "255A1B"
HAF_LIGHT_GREEN = "EFF5EC"
HAF_ROW_FILL = "C9D1C8"
HAF_ROW_ALT_FILL = "DCE3DA"
HAF_YELLOW = "E1D601"
HAF_TEXT = "1D1D1B"
WHITE = "FFFFFF"
OPINION_TABLE_COLUMNS = (
    ("company_name_ko", "종목명"),
    ("ticker", "티커"),
    ("market", "시장"),
    ("opinion", "의견"),
    ("score", "점수"),
    ("confidence", "신뢰도"),
    ("risk_level", "위험도"),
)
FINANCIAL_TABLE_ORDER = (
    ("per", "PER"),
    ("pbr", "PBR"),
    ("pcr", "PCR"),
    ("ev_to_ebitda", "EV/EBITDA"),
    ("eps", "EPS"),
    ("bps", "BPS"),
    ("ebitda", "EBITDA"),
    ("cash_dps", "CashDPS"),
    ("cash_dividend_yield", "Cash Dividend Yield"),
    ("roe", "ROE"),
    ("debt_ratio", "부채비율"),
    ("current_ratio", "유동비율"),
    ("operating_margin", "영업이익률"),
    ("net_margin", "순이익률"),
    ("market_cap", "시가총액"),
    ("enterprise_value", "기업가치(EV)"),
    ("shares_outstanding", "발행주식수"),
)
PERCENT_FIELDS = {
    "cash_dividend_yield",
    "roe",
    "operating_margin",
    "net_margin",
}
LARGE_NUMBER_FIELDS = {
    "ebitda",
    "market_cap",
    "enterprise_value",
    "shares_outstanding",
}


class WordReportError(RuntimeError):
    """Raised when a Word report cannot be generated."""


def fetch_selection_word_report(
    url: str,
    provider: str = DEFAULT_PROVIDER,
    model: Optional[str] = None,
    max_companies: int = 3,
    article_char_limit: int = DEFAULT_ARTICLE_CHAR_LIMIT,
    period: str = DEFAULT_PERIOD,
    interval: str = DEFAULT_INTERVAL,
    recent_rows: int = DEFAULT_RECENT_ROWS,
    output_dir: str = DEFAULT_WORD_REPORT_OUTPUT_DIR,
    recent_points: int = DEFAULT_RECENT_POINTS,
    daily_plot_period: str = DEFAULT_DAILY_PLOT_PERIOD,
    daily_plot_interval: str = DEFAULT_DAILY_PLOT_INTERVAL,
    daily_recent_points: int = DEFAULT_DAILY_RECENT_POINTS,
    plot_period: str = DEFAULT_PLOT_PERIOD,
    plot_interval: str = DEFAULT_PLOT_INTERVAL,
    clear_output_dir: bool = False,
    template_path: Optional[str] = None,
) -> dict[str, Any]:
    ensure_python_docx_available()

    research_result = fetch_selection_research_report(
        url=url,
        provider=provider,
        model=model,
        max_companies=max_companies,
        article_char_limit=article_char_limit,
        period=period,
        interval=interval,
        recent_rows=recent_rows,
    )

    output_path = Path(output_dir).resolve()
    output_path.mkdir(parents=True, exist_ok=True)
    plot_output_dir = output_path / "plots"
    plot_output_dir.mkdir(parents=True, exist_ok=True)

    daily_plot_companies = fetch_market_data_for_companies(
        companies=research_result["selection"].get("companies", []),
        period=daily_plot_period,
        interval=daily_plot_interval,
    )
    intraday_plot_companies = fetch_market_data_for_companies(
        companies=research_result["selection"].get("companies", []),
        period=plot_period,
        interval=plot_interval,
    )
    plots = generate_report_price_plots(
        daily_companies=daily_plot_companies,
        intraday_companies=intraday_plot_companies,
        output_dir=str(plot_output_dir),
        daily_recent_points=daily_recent_points,
        intraday_recent_points=recent_points,
        clear_output_dir=clear_output_dir,
    )
    chart_analysis_model = model if str(provider).lower() == "openai" else None
    plots = add_chart_analysis_to_plots(plots, model=chart_analysis_model)

    document_path = build_document_path(
        output_dir=output_path,
        article=research_result["article"],
    )
    resolved_template_path = resolve_word_report_template_path(template_path)
    generate_word_report_document(
        document_path=document_path,
        research_result=research_result,
        plots=plots,
        template_path=resolved_template_path,
    )

    return {
        **research_result,
        "plots": {
            "output_dir": str(plot_output_dir),
            "daily_period": daily_plot_period,
            "daily_interval": daily_plot_interval,
            "intraday_period": plot_period,
            "intraday_interval": plot_interval,
            "recent_points": recent_points,
            "daily_recent_points": daily_recent_points,
            "clear_output_dir": clear_output_dir,
            "companies": plots,
        },
        "word_report": {
            "document_path": str(document_path),
            "template_path": str(resolved_template_path) if resolved_template_path else None,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        },
    }


def generate_word_report_document(
    document_path: Path,
    research_result: dict[str, Any],
    plots: list[dict[str, Any]],
    template_path: Optional[Path] = None,
) -> None:
    ensure_python_docx_available()

    title, subtitle, sections = parse_markdown_report(
        research_result["llm_report"]["body"]
    )
    plot_map = {
        str(item.get("ticker", "")).strip(): item
        for item in plots
    }
    fundamentals_map = {
        str(item.get("ticker", "")).strip(): item
        for item in research_result["fundamentals"]["companies"]
    }
    technical_map = {
        str(item.get("ticker", "")).strip(): item
        for item in research_result["technical_analysis"]["companies"]
    }
    opinions_map = {
        str(item.get("ticker", "")).strip(): item
        for item in research_result["opinions"]["companies"]
    }

    company_order = [
        {
            "ticker": str(company.get("ticker", "")).strip(),
            "company_name_ko": str(company.get("company_name_ko", "")).strip(),
            "company_name": str(company.get("company_name", "")).strip(),
            "market": str(company.get("market", "")).strip(),
        }
        for company in research_result["selection"].get("companies", [])
    ]

    using_template = template_path is not None
    document = Document(str(template_path)) if using_template else Document()
    setattr(document, "_cam_template_mode", using_template)
    if not using_template:
        configure_document_styles(document)

    trailing_notice_table = detach_trailing_notice_table(document) if using_template else None
    if using_template:
        populate_word_report_template(
            document=document,
            research_result=research_result,
            plots=plots,
            title=title,
            subtitle=subtitle,
            company_order=company_order,
            fundamentals_map=fundamentals_map,
            technical_map=technical_map,
            opinions_map=opinions_map,
        )
        add_page_start(document)
        add_section_heading(document, "상세 리서치 본문", use_template_styles=True)
    section_lookup = {section["heading"]: section["blocks"] for section in sections}
    if not using_template:
        populate_haf_word_report(
            document=document,
            research_result=research_result,
            plots=plots,
            title=title,
            subtitle=subtitle,
            company_order=company_order,
            fundamentals_map=fundamentals_map,
            technical_map=technical_map,
            opinions_map=opinions_map,
            section_lookup=section_lookup,
        )
        if trailing_notice_table is not None:
            document._body._element.append(trailing_notice_table)
        document.save(document_path)
        return

    ordered_headings = [
        "Summary",
        "I. Investment Thesis",
        "II. Theme / Industry Interpretation",
        "III. Company-by-Company Analysis",
        "IV. Financial and Valuation Interpretation",
        "V. Technical / Momentum Interpretation",
        "VI. Risk Factors",
        "VII. Conclusion",
    ]

    inserted_company_exhibits = False
    for heading in ordered_headings:
        blocks = section_lookup.get(heading)
        if blocks is None:
            continue

        add_section_heading(document, heading, use_template_styles=using_template)
        add_markdown_blocks(document, blocks)

        if heading == "III. Company-by-Company Analysis":
            insert_company_exhibits(
                document=document,
                company_order=company_order,
                plot_map=plot_map,
                fundamentals_map=fundamentals_map,
                technical_map=technical_map,
                opinions_map=opinions_map,
                use_template_styles=using_template,
            )
            inserted_company_exhibits = True

    if not inserted_company_exhibits and company_order:
        add_section_heading(document, "기업별 차트 및 재무지표", use_template_styles=using_template)
        insert_company_exhibits(
            document=document,
            company_order=company_order,
            plot_map=plot_map,
            fundamentals_map=fundamentals_map,
            technical_map=technical_map,
            opinions_map=opinions_map,
            use_template_styles=using_template,
        )

    if trailing_notice_table is not None:
        document._body._element.append(trailing_notice_table)

    document.save(document_path)


def populate_haf_word_report(
    document: Any,
    research_result: dict[str, Any],
    plots: list[dict[str, Any]],
    title: str,
    subtitle: str,
    company_order: list[dict[str, str]],
    fundamentals_map: dict[str, dict[str, Any]],
    technical_map: dict[str, dict[str, Any]],
    opinions_map: dict[str, dict[str, Any]],
    section_lookup: dict[str, list[dict[str, Any]]],
) -> None:
    article = research_result["article"]
    template_date = format_template_date(article.get("published_at"))
    plot_map = {
        str(item.get("ticker", "")).strip(): item
        for item in plots
    }

    configure_haf_section_header(document.sections[0], template_date)
    add_haf_cover_page(
        document=document,
        research_result=research_result,
        title=title,
        subtitle=subtitle,
        company_order=company_order,
        opinions_map=opinions_map,
        section_lookup=section_lookup,
    )

    add_page_start(document)
    add_haf_overview_page(
        document=document,
        section_lookup=section_lookup,
        company_order=company_order,
        opinions_map=opinions_map,
    )
    add_page_start(document)
    add_haf_interpretation_page(
        document=document,
        section_lookup=section_lookup,
    )

    for company in company_order:
        add_page_start(document)
        add_haf_company_overview_page(
            document=document,
            company=company,
            fundamentals=fundamentals_map.get(company["ticker"], {}),
            technical=technical_map.get(company["ticker"], {}),
            opinion=opinions_map.get(company["ticker"], {}),
            company_analysis_blocks=section_lookup.get("III. Company-by-Company Analysis", []),
            risk_blocks=section_lookup.get("VI. Risk Factors", []),
        )
        add_page_start(document)
        add_haf_company_exhibit_page(
            document=document,
            company=company,
            plot=plot_map.get(company["ticker"], {}),
            fundamentals=fundamentals_map.get(company["ticker"], {}),
            technical=technical_map.get(company["ticker"], {}),
            opinion=opinions_map.get(company["ticker"], {}),
        )


def add_haf_cover_page(
    document: Any,
    research_result: dict[str, Any],
    title: str,
    subtitle: str,
    company_order: list[dict[str, str]],
    opinions_map: dict[str, dict[str, Any]],
    section_lookup: dict[str, list[dict[str, Any]]],
) -> None:
    article = research_result["article"]
    title_text = title or article.get("title") or "리서치 리포트"
    subtitle_text = (
        subtitle
        or str(research_result["selection"].get("summary", "")).strip()
        or "헤더 기반 워드 리포트"
    )

    add_horizontal_rule(document, HAF_GREEN, 12)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(6)
    add_styled_run(
        title_paragraph,
        title_text,
        size=22,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_TEXT,
    )

    accent_bar = document.add_table(rows=1, cols=2)
    accent_bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    accent_bar.autofit = False
    accent_bar.columns[0].width = Inches(5.8)
    accent_bar.columns[1].width = Inches(2.0)
    left_bar = accent_bar.cell(0, 0)
    right_bar = accent_bar.cell(0, 1)
    set_cell_background(left_bar, HAF_GREEN)
    set_cell_background(right_bar, HAF_YELLOW)
    set_cell_text(left_bar, "", size=1, font_name=None)
    set_cell_text(right_bar, "", size=1, font_name=None)
    remove_cell_borders(left_bar)
    remove_cell_borders(right_bar)
    set_cell_margins(left_bar, top=0, bottom=0, start=0, end=0)
    set_cell_margins(right_bar, top=0, bottom=0, start=0, end=0)

    hero_box = document.add_table(rows=1, cols=1)
    hero_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    hero_cell = hero_box.cell(0, 0)
    set_cell_background(hero_cell, HAF_LIGHT_GREEN)
    remove_cell_borders(hero_cell)
    set_cell_margins(hero_cell, top=180, bottom=180, start=220, end=220)
    set_cell_text(
        hero_cell,
        subtitle_text,
        bold=True,
        size=18,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_DEEP_GREEN,
    )

    summary_table = build_haf_summary_table(document, company_order, opinions_map)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    summary_text = blocks_to_plain_text(section_lookup.get("Summary", []), max_chars=420)
    industry_text = blocks_to_plain_text(
        section_lookup.get("I. Investment Thesis", [])
        + section_lookup.get("II. Theme / Industry Interpretation", []),
        max_chars=780,
    )

    add_haf_text_card(document, "Summary", summary_text)
    add_haf_text_card(document, "INDUSTRY", industry_text)


def add_haf_overview_page(
    document: Any,
    section_lookup: dict[str, list[dict[str, Any]]],
    company_order: list[dict[str, str]],
    opinions_map: dict[str, dict[str, Any]],
) -> None:
    conclusion_text = blocks_to_plain_text(
        section_lookup.get("VII. Conclusion", []),
        max_chars=900,
    )
    risk_text = blocks_to_plain_text(
        section_lookup.get("VI. Risk Factors", []),
        max_chars=700,
    )

    add_haf_text_card(document, "Conclusion", conclusion_text)
    add_haf_text_card(document, "Key Factors", build_key_factor_text(company_order, opinions_map, risk_text))


def add_haf_interpretation_page(
    document: Any,
    section_lookup: dict[str, list[dict[str, Any]]],
) -> None:
    valuation_text = blocks_to_plain_text(
        section_lookup.get("IV. Financial and Valuation Interpretation", []),
        max_chars=1200,
    )
    technical_text = blocks_to_plain_text(
        section_lookup.get("V. Technical / Momentum Interpretation", []),
        max_chars=1200,
    )
    risk_text = blocks_to_plain_text(
        section_lookup.get("VI. Risk Factors", []),
        max_chars=1000,
    )

    add_haf_text_card(document, "Financial / Valuation", valuation_text)
    add_haf_text_card(document, "Technical / Momentum", technical_text)
    add_haf_text_card(document, "Risk Factors", risk_text)


def add_haf_company_overview_page(
    document: Any,
    company: dict[str, str],
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
    opinion: dict[str, Any],
    company_analysis_blocks: list[dict[str, Any]],
    risk_blocks: list[dict[str, Any]],
) -> None:
    layout = document.add_table(rows=1, cols=2)
    layout.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout.autofit = False
    layout.columns[0].width = Inches(2.3)
    layout.columns[1].width = Inches(4.8)
    left_cell = layout.cell(0, 0)
    right_cell = layout.cell(0, 1)
    set_cell_background(left_cell, HAF_GREEN)
    remove_cell_borders(left_cell)
    remove_cell_borders(right_cell)
    set_cell_margins(left_cell, top=160, bottom=180, start=160, end=160)
    set_cell_margins(right_cell, top=60, bottom=80, start=150, end=150)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    left_title = left_cell.paragraphs[0]
    add_styled_run(
        left_title,
        f"{company.get('company_name_ko') or company.get('company_name')}\n({company.get('ticker')})",
        size=18,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=WHITE,
    )

    left_meta = left_cell.add_paragraph()
    add_styled_run(
        left_meta,
        (
            f"{company.get('market') or 'N/A'}\n"
            f"재무 기준일: {fundamentals.get('statement_date') or 'N/A'}\n"
            f"출처: {fundamentals.get('source') or 'OpenDART'}"
        ),
        size=9.5,
        font_name=DEFAULT_BODY_FONT,
        color=WHITE,
    )

    metrics_box = left_cell.add_table(rows=1, cols=1)
    metrics_cell = metrics_box.cell(0, 0)
    set_cell_background(metrics_cell, WHITE)
    remove_cell_borders(metrics_cell)
    set_cell_margins(metrics_cell, top=100, bottom=100, start=110, end=110)
    set_cell_text(
        metrics_cell,
        build_haf_metric_box(opinion, fundamentals),
        size=9,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_TEXT,
    )

    opinion_text = build_company_opinion_text(company_analysis_blocks, company, opinion, fundamentals)
    technical_text = build_company_technical_text(technical, opinion, fundamentals)
    risk_text = build_company_risk_text(
        risk_blocks=risk_blocks,
        company=company,
        opinion=opinion,
        fundamentals=fundamentals,
        technical=technical,
    )

    add_haf_text_card(right_cell, "OPINION", opinion_text, compact=True)
    add_haf_text_card(right_cell, "Technical / Momentum", technical_text, compact=True)
    add_haf_text_card(right_cell, "RISK", risk_text, compact=True)


def add_haf_company_exhibit_page(
    document: Any,
    company: dict[str, str],
    plot: dict[str, Any],
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
    opinion: dict[str, Any],
) -> None:
    header = document.add_paragraph()
    add_styled_run(
        header,
        f"{company.get('company_name_ko') or company.get('company_name')} ({company.get('ticker')})",
        size=18,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_DEEP_GREEN,
    )

    subheader = document.add_paragraph()
    add_styled_run(
        subheader,
        (
            f"{company.get('market') or 'N/A'} | "
            f"의견 {str(opinion.get('opinion') or 'N/A')} | "
            f"점수 {format_scalar(opinion.get('score'), decimals=2)} | "
            f"신뢰도 {format_scalar(opinion.get('confidence'), decimals=2)}"
        ),
        size=10,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_TEXT,
    )

    daily_plot_path = str(plot.get("daily_plot_path") or plot.get("plot_path") or "").strip()
    intraday_plot_path = str(plot.get("intraday_plot_path") or "").strip()
    if daily_plot_path and Path(daily_plot_path).exists():
        add_haf_large_plot_card(
            document=document,
            title=f"{company.get('company_name_ko') or company.get('ticker')} 3M Daily Price Trend",
            plot_path=daily_plot_path,
            caption=build_plot_caption(plot, chart_kind="daily"),
        )
        add_chart_comment_card(
            document,
            "LLM Technical Comment - 3M Daily",
            plot.get("daily_chart_comment"),
        )
    if intraday_plot_path and Path(intraday_plot_path).exists():
        add_haf_large_plot_card(
            document=document,
            title=f"{company.get('company_name_ko') or company.get('ticker')} 1M Intraday Price / Volume",
            plot_path=intraday_plot_path,
            caption=build_plot_caption(plot, chart_kind="intraday"),
        )
        add_chart_comment_card(
            document,
            "LLM Technical Comment - 1M Intraday",
            plot.get("intraday_chart_comment"),
        )

    add_haf_financial_snapshot_page(
        document=document,
        fundamentals=fundamentals,
        technical=technical,
    )


def configure_haf_section_header(section: Any, template_date: str) -> None:
    section.header_distance = Inches(0.22)
    header = section.header
    header.is_linked_to_previous = False

    if header.paragraphs:
        header.paragraphs[0].text = ""

    header_table = header.add_table(rows=1, cols=2, width=usable_section_width(section))
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(5.4)
    header_table.columns[1].width = Inches(1.8)
    left_cell = header_table.cell(0, 0)
    right_cell = header_table.cell(0, 1)
    remove_cell_borders(left_cell)
    remove_cell_borders(right_cell)
    set_cell_border(left_cell, bottom={"color": HAF_GREEN, "sz": 12})
    set_cell_border(right_cell, bottom={"color": HAF_GREEN, "sz": 12})
    set_cell_margins(left_cell, top=0, bottom=40, start=0, end=0)
    set_cell_margins(right_cell, top=0, bottom=40, start=0, end=0)

    set_cell_text(
        left_cell,
        f"Research | {template_date}",
        size=11,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_TEXT,
    )
    set_cell_text(
        right_cell,
        "HAF",
        size=17,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_GREEN,
    )
    right_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_haf_summary_table(
    document: Any,
    company_order: list[dict[str, str]],
    opinions_map: dict[str, dict[str, Any]],
) -> Any:
    table = document.add_table(rows=1, cols=len(OPINION_TABLE_COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for index, (_, label) in enumerate(OPINION_TABLE_COLUMNS):
        cell = table.rows[0].cells[index]
        set_cell_background(cell, HAF_GREEN)
        set_cell_text(
            cell,
            label,
            bold=True,
            size=10,
            font_name=DEFAULT_BODY_FONT,
            color=WHITE,
        )
        set_cell_margins(cell, top=70, bottom=70, start=80, end=80)

    for row_index, company in enumerate(company_order[:3]):
        row_cells = table.add_row().cells
        opinion = opinions_map.get(company["ticker"], {})
        fill = HAF_ROW_FILL if row_index % 2 == 0 else HAF_ROW_ALT_FILL
        values = (
            company.get("company_name_ko") or company.get("company_name") or "N/A",
            company.get("ticker") or "N/A",
            company.get("market") or "N/A",
            str(opinion.get("opinion") or "N/A"),
            format_scalar(opinion.get("score"), decimals=2),
            format_scalar(opinion.get("confidence"), decimals=2),
            str(opinion.get("risk_level") or "N/A"),
        )
        for idx, value in enumerate(values):
            cell = row_cells[idx]
            set_cell_background(cell, fill)
            set_cell_text(
                cell,
                value,
                size=9.5,
                font_name=DEFAULT_BODY_FONT,
                color=HAF_TEXT,
            )
            set_cell_margins(cell, top=70, bottom=70, start=80, end=80)

    return table


def add_haf_text_card(
    container: Any,
    title: str,
    body: str,
    compact: bool = False,
) -> None:
    card = container.add_table(rows=2, cols=1)
    card.alignment = WD_TABLE_ALIGNMENT.CENTER
    card.autofit = True
    title_cell = card.cell(0, 0)
    body_cell = card.cell(1, 0)

    set_cell_background(title_cell, WHITE)
    set_cell_background(body_cell, HAF_LIGHT_GREEN if not compact else WHITE)
    remove_cell_borders(title_cell)
    remove_cell_borders(body_cell)
    set_cell_border(title_cell, bottom={"color": HAF_GREEN, "sz": 6})
    set_cell_margins(title_cell, top=40, bottom=80, start=0, end=0)
    set_cell_margins(body_cell, top=120 if not compact else 90, bottom=120, start=120, end=120)

    set_cell_text(
        title_cell,
        title,
        bold=True,
        size=17 if not compact else 13,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_DEEP_GREEN,
    )
    set_cell_text(
        body_cell,
        body or "내용이 없습니다.",
        size=11 if not compact else 10.2,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_TEXT,
    )
    body_cell.paragraphs[0].paragraph_format.space_after = Pt(0)


def add_haf_financial_card(
    container: Any,
    fundamentals: dict[str, Any],
    title_text: str = "Financial Indicators",
) -> None:
    title = container.add_paragraph()
    title.paragraph_format.space_before = Pt(4)
    add_styled_run(
        title,
        title_text,
        size=13,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_DEEP_GREEN,
    )
    insert_financial_analysis_table(container, fundamentals)


def add_haf_financial_snapshot_page(
    document: Any,
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
) -> None:
    metrics = fundamentals.get("valuation_metrics", {})
    highlight_table = document.add_table(rows=2, cols=4)
    highlight_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    highlight_table.autofit = False

    highlight_items = (
        ("종가", format_financial_metric("close", fundamentals.get("latest_close"))),
        ("PER", format_financial_metric("per", metrics.get("per"))),
        ("PBR", format_financial_metric("pbr", metrics.get("pbr"))),
        ("ROE", format_financial_metric("roe", metrics.get("roe"))),
    )

    for index, (label, value) in enumerate(highlight_items):
        label_cell = highlight_table.rows[0].cells[index]
        value_cell = highlight_table.rows[1].cells[index]
        set_cell_background(label_cell, HAF_GREEN)
        set_cell_background(value_cell, HAF_LIGHT_GREEN)
        set_cell_text(
            label_cell,
            label,
            bold=True,
            size=10,
            font_name=DEFAULT_BODY_FONT,
            color=WHITE,
        )
        set_cell_text(
            value_cell,
            value,
            size=10.5,
            font_name=DEFAULT_BODY_FONT,
            color=HAF_TEXT,
        )
        set_cell_margins(label_cell, top=70, bottom=70, start=80, end=80)
        set_cell_margins(value_cell, top=90, bottom=90, start=80, end=80)

    add_haf_text_card(
        document,
        "Financial Snapshot",
        build_financial_snapshot_text(fundamentals, technical),
    )
    add_haf_financial_card(document, fundamentals)


def add_haf_large_plot_card(
    document: Any,
    title: str,
    plot_path: str,
    caption: str,
) -> None:
    card = document.add_table(rows=2, cols=1)
    card.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_cell = card.cell(0, 0)
    body_cell = card.cell(1, 0)

    remove_cell_borders(title_cell)
    remove_cell_borders(body_cell)
    set_cell_border(title_cell, bottom={"color": HAF_GREEN, "sz": 6})
    set_cell_margins(title_cell, top=60, bottom=80, start=0, end=0)
    set_cell_margins(body_cell, top=120, bottom=90, start=90, end=90)

    set_cell_text(
        title_cell,
        title,
        bold=True,
        size=13,
        font_name=DEFAULT_BODY_FONT,
        color=HAF_DEEP_GREEN,
    )

    body_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    body_cell.paragraphs[0].add_run().add_picture(plot_path, width=Inches(6.35))
    caption_paragraph = body_cell.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_styled_run(
        caption_paragraph,
        caption,
        size=9,
        italic=True,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_SOFT_GRAY,
    )


def add_chart_comment_card(
    document: Any,
    title: str,
    comment: Any,
) -> None:
    text = str(comment or "").strip()
    if not text:
        return
    add_haf_text_card(
        document,
        title,
        text,
        compact=True,
    )


def add_horizontal_rule(container: Any, color: str, size: int) -> None:
    rule = container.add_table(rows=1, cols=1)
    cell = rule.cell(0, 0)
    set_cell_text(cell, "", size=1, font_name=None)
    remove_cell_borders(cell, keep_bottom=True)
    set_cell_border(cell, bottom={"color": color, "sz": size})
    set_cell_margins(cell, top=0, bottom=0, start=0, end=0)


def add_page_start(document: Any) -> None:
    anchor = document.add_paragraph()
    anchor.paragraph_format.page_break_before = True
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(0)
    if anchor.runs:
        anchor.runs[0].font.hidden = True


def build_key_factor_text(
    company_order: list[dict[str, str]],
    opinions_map: dict[str, dict[str, Any]],
    risk_text: str,
) -> str:
    lines = []
    for company in company_order[:3]:
        opinion = opinions_map.get(company["ticker"], {})
        score = format_scalar(opinion.get("score"), decimals=2)
        confidence = format_scalar(opinion.get("confidence"), decimals=2)
        lines.append(
            f"{company.get('company_name_ko') or company.get('ticker')}: "
            f"{str(opinion.get('opinion') or 'N/A')} / score {score} / confidence {confidence}"
        )

    if risk_text:
        lines.append(f"공통 리스크: {risk_text}")
    return "\n".join(lines)


def build_company_opinion_text(
    company_analysis_blocks: list[dict[str, Any]],
    company: dict[str, str],
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
) -> str:
    llm_text = extract_company_reference_text(company_analysis_blocks, company, max_chars=420)
    if llm_text:
        return llm_text

    pieces = []
    rationale = str(opinion.get("rationale") or "").strip()
    if rationale:
        pieces.append(rationale)
    latest_close = format_financial_metric("close", fundamentals.get("latest_close"))
    if latest_close != "N/A":
        pieces.append(f"최신 종가는 {latest_close}이다.")
    positives = ", ".join(str(item).strip() for item in opinion.get("positives", []) if str(item).strip())
    if positives:
        pieces.append(f"긍정 요인은 {positives}이다.")
    return compact_text(" ".join(pieces), max_chars=420)


def build_company_technical_text(
    technical: dict[str, Any],
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
) -> str:
    latest = technical.get("latest_indicators", {})
    latest_signal = technical.get("latest_signal", {})
    if not isinstance(latest_signal, dict):
        latest_signal = {}

    signal = str(latest_signal.get("signal") or "").strip()
    signal_label = translate_trade_signal(signal)
    regime = str(
        technical.get("signal_context", {}).get("market_regime", "")
        or latest_signal.get("regime", "")
    ).strip()
    regime_label = translate_market_regime(regime)
    sentences = []
    score = format_scalar(latest_signal.get("score", opinion.get("score")), decimals=2)
    confidence = format_scalar(opinion.get("confidence"), decimals=2)

    if regime_label != "미확인" or signal_label != "분석 불가":
        lead = "현재 기술 국면은 "
        if regime_label != "미확인":
            lead += f"{regime_label}으로 분류되며, "
        else:
            lead += "명확히 확정되진 않았으나, "
        if signal_label != "분석 불가":
            lead += f"종합 신호는 {signal_label}"
        else:
            lead += "종합 신호는 해석 제한"
        if score != "N/A":
            lead += f"(점수 {score}"
            if confidence != "N/A":
                lead += f", 신뢰도 {confidence}"
            lead += ")"
        lead += "다."
        sentences.append(lead)

    trend_sentence = build_trend_structure_sentence(latest, fundamentals)
    if trend_sentence:
        sentences.append(trend_sentence)

    momentum_sentence = build_momentum_structure_sentence(latest, technical.get("signal_context", {}))
    if momentum_sentence:
        sentences.append(momentum_sentence)

    oscillator_sentence = build_oscillator_structure_sentence(latest, technical.get("signal_context", {}))
    if oscillator_sentence:
        sentences.append(oscillator_sentence)

    evidence_sentence = build_breakdown_evidence_sentence(opinion)
    if evidence_sentence:
        sentences.append(evidence_sentence)

    if not sentences:
        return "기술적 해석에 필요한 지표가 충분하지 않습니다."

    return compact_text(" ".join(sentences), max_chars=780)


def build_company_risk_text(
    risk_blocks: list[dict[str, Any]],
    company: dict[str, str],
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
) -> str:
    company_risk_text = extract_company_reference_text(risk_blocks, company, max_chars=260)
    reasons = build_company_risk_reasons(
        company_risk_text=company_risk_text,
        opinion=opinion,
        fundamentals=fundamentals,
        technical=technical,
    )

    risk_level = str(opinion.get("risk_level") or "").strip()
    intro = f"위험도는 {risk_level} 수준으로 판단된다. " if risk_level else ""
    numbered = []
    for index, reason in enumerate(reasons[:4], start=1):
        numbered.append(f"{ordinal_korean(index)}, {reason}")
    return compact_text(intro + " ".join(numbered), max_chars=620)


def build_haf_metric_box(
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
) -> str:
    metrics = fundamentals.get("valuation_metrics", {})
    rows = [
        f"의견      {str(opinion.get('opinion') or 'N/A')}",
        f"점수      {format_scalar(opinion.get('score'), decimals=2)}",
        f"신뢰도    {format_scalar(opinion.get('confidence'), decimals=2)}",
        f"종가      {format_financial_metric('close', fundamentals.get('latest_close'))}",
        f"PER       {format_financial_metric('per', metrics.get('per'))}",
        f"PBR       {format_financial_metric('pbr', metrics.get('pbr'))}",
        f"ROE       {format_financial_metric('roe', metrics.get('roe'))}",
    ]
    return "\n".join(rows)


def build_plot_caption(plot: dict[str, Any], chart_kind: str = "daily") -> str:
    if chart_kind == "intraday":
        start_date = str(plot.get("intraday_start_date") or "").strip()
        end_date = str(plot.get("intraday_end_date") or "").strip()
        latest_close = format_financial_metric("close", plot.get("intraday_latest_close"))
        label = "1분봉 주가/거래량"
    else:
        start_date = str(plot.get("daily_start_date") or plot.get("start_date") or "").strip()
        end_date = str(plot.get("daily_end_date") or plot.get("end_date") or "").strip()
        latest_close = format_financial_metric("close", plot.get("daily_latest_close") or plot.get("latest_close"))
        label = "3개월 일봉 종가 흐름"

    if start_date and end_date and latest_close != "N/A":
        return f"{label} ({start_date} ~ {end_date}) | Latest close {latest_close}"
    if start_date and end_date:
        return f"{label} ({start_date} ~ {end_date})"
    return label


def build_financial_snapshot_text(
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
) -> str:
    metrics = fundamentals.get("valuation_metrics", {})
    parts: list[str] = []

    statement_date = str(fundamentals.get("statement_date") or "").strip()
    if statement_date:
        parts.append(f"재무 기준일은 {statement_date}이다.")

    per_value = format_financial_metric("per", metrics.get("per"))
    pbr_value = format_financial_metric("pbr", metrics.get("pbr"))
    roe_value = format_financial_metric("roe", metrics.get("roe"))
    market_cap_value = format_financial_metric("market_cap", metrics.get("market_cap"))

    valuation_fragments = []
    if per_value != "N/A":
        valuation_fragments.append(f"PER은 {per_value}")
    if pbr_value != "N/A":
        valuation_fragments.append(f"PBR은 {pbr_value}")
    if roe_value != "N/A":
        valuation_fragments.append(f"ROE는 {roe_value}")
    if valuation_fragments:
        parts.append(", ".join(valuation_fragments) + " 수준이다.")

    if market_cap_value != "N/A":
        parts.append(f"시가총액은 {market_cap_value}로 집계된다.")

    regime = translate_market_regime(
        str(technical.get("signal_context", {}).get("market_regime", "")).strip()
    )
    if regime != "미확인":
        parts.append(f"기술적 시장 국면은 {regime}으로 분류된다.")

    if not parts:
        return "핵심 재무 요약이 없습니다."

    return compact_text(" ".join(parts), max_chars=520)


def build_company_risk_reasons(
    company_risk_text: str,
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
) -> list[str]:
    reasons: list[str] = []
    negatives = [
        str(item).strip()
        for item in opinion.get("negatives", [])
        if str(item).strip()
    ]
    for negative in negatives:
        append_unique_reason(reasons, f"{negative} 요인이 주가 변동성을 키울 수 있다.")

    for sentence in split_text_into_sentences(company_risk_text):
        cleaned = sentence.strip().rstrip(".")
        if cleaned:
            append_unique_reason(reasons, cleaned + ".")

    metrics = fundamentals.get("valuation_metrics", {})
    per_value = try_float(metrics.get("per"))
    pbr_value = try_float(metrics.get("pbr"))
    if per_value is not None and per_value >= 35:
        append_unique_reason(
            reasons,
            f"PER이 {format_financial_metric('per', per_value)} 수준으로 높아 기대가 이미 가격에 반영됐을 가능성이 있어 밸류에이션 조정 부담이 있다.",
        )
    if pbr_value is not None and pbr_value >= 3:
        append_unique_reason(
            reasons,
            f"PBR이 {format_financial_metric('pbr', pbr_value)} 수준으로 높아 실적 둔화 시 멀티플 축소 압력이 커질 수 있다.",
        )

    latest = technical.get("latest_indicators", {})
    rsi_value = try_float(latest.get("rsi_14"))
    if rsi_value is not None and rsi_value >= 70:
        append_unique_reason(
            reasons,
            f"RSI가 {format_scalar(rsi_value, decimals=2)}로 과매수 구간에 있어 단기 차익실현 압력이 확대될 수 있다.",
        )
    elif rsi_value is not None and rsi_value <= 30:
        append_unique_reason(
            reasons,
            f"RSI가 {format_scalar(rsi_value, decimals=2)}로 과매도 구간에 머물러 변동성이 커질 수 있다.",
        )

    adx_value = try_float(latest.get("adx_14"))
    if adx_value is not None and adx_value < 23:
        append_unique_reason(
            reasons,
            f"ADX가 {format_scalar(adx_value, decimals=2)}로 높지 않아 추세 지속성이 약해지고 박스권 잡음이 커질 수 있다.",
        )

    obv_trend = str(technical.get("signal_context", {}).get("obv_trend", "")).strip()
    if obv_trend == "bearish":
        append_unique_reason(
            reasons,
            "OBV 흐름이 약세로 분류돼 거래량이 상승 흐름을 충분히 뒷받침하지 못할 가능성이 있다.",
        )

    confidence = try_float(opinion.get("confidence"))
    if confidence is not None and confidence < 0.75:
        append_unique_reason(
            reasons,
            f"신뢰도가 {format_scalar(confidence, decimals=2)} 수준으로 아주 높지는 않아 기술 신호의 일관성이 약해질 가능성도 염두에 둘 필요가 있다.",
        )

    fallback_reasons = [
        "테마 기대감이 빠르게 약해질 경우 수급 방향이 급격히 바뀔 수 있다.",
        "실적 확인 전까지는 뉴스 모멘텀에 따라 주가가 과민하게 반응할 가능성이 있다.",
        "시장 전반의 위험회피 심리가 커지면 개별 종목의 상대 강도도 빠르게 약해질 수 있다.",
    ]
    for fallback in fallback_reasons:
        if len(reasons) >= 3:
            break
        append_unique_reason(reasons, fallback)

    return reasons[:4]


def append_unique_reason(reasons: list[str], reason: str) -> None:
    cleaned = " ".join(str(reason or "").split()).strip()
    if not cleaned:
        return
    if cleaned not in reasons:
        reasons.append(cleaned)


def ordinal_korean(index: int) -> str:
    if index == 1:
        return "첫째"
    if index == 2:
        return "둘째"
    if index == 3:
        return "셋째"
    if index == 4:
        return "넷째"
    return f"{index}번째"


def build_trend_structure_sentence(
    latest: dict[str, Any],
    fundamentals: dict[str, Any],
) -> str:
    close_value = latest.get("close")
    if close_value is None:
        close_value = fundamentals.get("latest_close")
    sma20 = latest.get("sma_20")
    sma50 = latest.get("sma_50")
    ema20 = latest.get("ema_20")

    parts: list[str] = []
    if sma20 is not None and sma50 is not None:
        if sma20 > sma50:
            parts.append(
                f"최근 20일 평균가격과 50일 평균가격의 상대 위치로 중기 추세를 읽는 이동평균 지표에서, SMA20({format_financial_metric('close', sma20)})이 SMA50({format_financial_metric('close', sma50)}) 위에 있어 중기 추세는 우상향으로 해석된다"
            )
        elif sma20 < sma50:
            parts.append(
                f"최근 20일 평균가격과 50일 평균가격의 상대 위치로 중기 추세를 읽는 이동평균 지표에서, SMA20({format_financial_metric('close', sma20)})이 SMA50({format_financial_metric('close', sma50)}) 아래에 있어 중기 흐름은 상대적으로 약한 편이다"
            )

    if close_value is not None and ema20 is not None:
        if close_value > ema20:
            parts.append(
                f"최근 가격 압력이 단기적으로 얼마나 강한지를 보여주는 EMA20 기준으로도 종가({format_financial_metric('close', close_value)})가 EMA20({format_financial_metric('close', ema20)}) 위에 있어 단기 매수 압력이 우세하다"
            )
        elif close_value < ema20:
            parts.append(
                f"최근 가격 압력이 단기적으로 얼마나 강한지를 보여주는 EMA20 기준으로는 종가({format_financial_metric('close', close_value)})가 EMA20({format_financial_metric('close', ema20)}) 아래에 있어 단기 탄력은 다소 제한적이다"
            )

    return ". ".join(parts) + ("." if parts else "")


def build_momentum_structure_sentence(
    latest: dict[str, Any],
    signal_context: dict[str, Any],
) -> str:
    macd = latest.get("macd")
    macd_signal = latest.get("macd_signal")
    macd_hist = latest.get("macd_histogram")
    adx = latest.get("adx_14")
    plus_di = latest.get("dmp_14")
    minus_di = latest.get("dmn_14")
    obv_trend = str(signal_context.get("obv_trend", "")).strip()

    parts: list[str] = []
    if macd is not None and macd_signal is not None:
        if macd > macd_signal:
            text = (
                f"단기·장기 이동평균의 간격 변화로 모멘텀 전환을 읽는 MACD 지표는 {format_scalar(macd, decimals=2)}이며, 시그널선({format_scalar(macd_signal, decimals=2)}) 위에 있어 상승 모멘텀이 우세하다"
            )
            if macd_hist is not None and macd_hist > 0:
                text += f"고, 히스토그램도 양(+)의 구간({format_scalar(macd_hist, decimals=2)})에 있다"
            parts.append(text)
        elif macd < macd_signal:
            text = (
                f"단기·장기 이동평균의 간격 변화로 모멘텀 전환을 읽는 MACD 지표는 {format_scalar(macd, decimals=2)}이며, 시그널선({format_scalar(macd_signal, decimals=2)}) 아래에 있어 하락 모멘텀이 우세하다"
            )
            if macd_hist is not None and macd_hist < 0:
                text += f"고, 히스토그램도 음(-)의 구간({format_scalar(macd_hist, decimals=2)})에 있다"
            parts.append(text)

    if adx is not None and plus_di is not None and minus_di is not None:
        trend_strength = "강한 편" if adx >= 30 else "보통" if adx >= 23 else "약한 편"
        if plus_di > minus_di:
            parts.append(
                f"추세의 세기 자체를 측정하는 ADX는 {format_scalar(adx, decimals=2)} 수준으로 {trend_strength}이고, 방향성 지표에서는 +DI({format_scalar(plus_di, decimals=2)})가 -DI({format_scalar(minus_di, decimals=2)})보다 높아 상승 추세가 우세하다"
            )
        elif plus_di < minus_di:
            parts.append(
                f"추세의 세기 자체를 측정하는 ADX는 {format_scalar(adx, decimals=2)} 수준으로 {trend_strength}이고, 방향성 지표에서는 -DI({format_scalar(minus_di, decimals=2)})가 +DI({format_scalar(plus_di, decimals=2)})보다 높아 하락 추세가 우세하다"
            )

    if obv_trend == "bullish":
        parts.append("거래량이 가격 움직임을 얼마나 뒷받침하는지 보여주는 OBV 흐름도 단기 평균이 장기 평균 위에 있어 수급 측면의 뒷받침이 우호적이다")
    elif obv_trend == "bearish":
        parts.append("거래량이 가격 움직임을 얼마나 뒷받침하는지 보여주는 OBV 흐름은 단기 평균이 장기 평균 아래에 있어 수급 측면의 뒷받침이 약하다")

    return ". ".join(parts) + ("." if parts else "")


def build_oscillator_structure_sentence(
    latest: dict[str, Any],
    signal_context: dict[str, Any],
) -> str:
    rsi = latest.get("rsi_14")
    stoch_k = latest.get("stoch_k")
    stoch_d = latest.get("stoch_d")
    bollinger_position = str(signal_context.get("bollinger_position", "")).strip()

    parts: list[str] = []
    if rsi is not None:
        if rsi >= 70:
            parts.append(f"가격의 상승·하락 압력이 얼마나 한쪽으로 쏠렸는지 보여주는 RSI 지표는 {format_scalar(rsi, decimals=2)}로, 통상적 기준상 과매수(70 이상) 구간에 해당한다")
        elif rsi <= 30:
            parts.append(f"가격의 상승·하락 압력이 얼마나 한쪽으로 쏠렸는지 보여주는 RSI 지표는 {format_scalar(rsi, decimals=2)}로, 통상적 기준상 과매도(30 이하) 구간에 해당한다")
        elif rsi >= 60:
            parts.append(f"가격의 상승·하락 압력이 강세 쪽으로 기울었는지 보는 RSI 지표는 {format_scalar(rsi, decimals=2)}로, 과열은 아니지만 강세권에 위치한다")
        elif rsi <= 40:
            parts.append(f"가격의 상승·하락 압력이 약세 쪽으로 기울었는지 보는 RSI 지표는 {format_scalar(rsi, decimals=2)}로, 약세권에 가깝다")
        else:
            parts.append(f"가격의 상승·하락 압력이 어느 한쪽으로 과도하게 쏠리지 않았는지 보는 RSI 지표는 {format_scalar(rsi, decimals=2)}로 중립권에 가깝다")

    if stoch_k is not None and stoch_d is not None:
        if stoch_k > stoch_d and stoch_k < 20:
            parts.append(
                f"단기 변곡 가능성을 빠르게 포착하는 스토캐스틱은 %K {format_scalar(stoch_k, decimals=2)}, %D {format_scalar(stoch_d, decimals=2)}로 과매도권 반등 신호에 가깝다"
            )
        elif stoch_k < stoch_d and stoch_k > 80:
            parts.append(
                f"단기 변곡 가능성을 빠르게 포착하는 스토캐스틱은 %K {format_scalar(stoch_k, decimals=2)}, %D {format_scalar(stoch_d, decimals=2)}로 과매수권 둔화 신호에 가깝다"
            )

    if bollinger_position == "above_upper_band":
        parts.append("변동성 범위 대비 현재 가격 위치를 보여주는 볼린저밴드 기준으로는 주가가 상단을 웃돌아 단기 과열 부담이 있다")
    elif bollinger_position == "below_lower_band":
        parts.append("변동성 범위 대비 현재 가격 위치를 보여주는 볼린저밴드 기준으로는 주가가 하단 아래에 있어 단기 과매도 성격이 있다")

    return ". ".join(parts) + ("." if parts else "")


def build_breakdown_evidence_sentence(opinion: dict[str, Any]) -> str:
    breakdown = opinion.get("score_breakdown", [])
    if not isinstance(breakdown, list):
        return ""

    ranked = []
    for item in breakdown:
        if not isinstance(item, dict):
            continue
        detail = str(item.get("detail", "")).strip()
        score = try_float(item.get("score"))
        if not detail or score is None or score == 0:
            continue
        ranked.append((abs(score), detail))

    if not ranked:
        return ""

    seen: list[str] = []
    for _, detail in sorted(ranked, key=lambda item: item[0], reverse=True):
        cleaned = detail.rstrip(". ").strip()
        if cleaned and cleaned not in seen:
            seen.append(cleaned)
        if len(seen) == 3:
            break

    return f"주요 근거로는 {', '.join(seen)} 등을 들 수 있다."


def translate_trade_signal(signal: str) -> str:
    normalized = str(signal or "").strip().upper()
    if normalized == "STRONG_BUY":
        return "강력 매수"
    if normalized == "BUY":
        return "매수"
    if normalized == "HOLD":
        return "보유"
    if normalized == "SELL":
        return "매도"
    if normalized == "STRONG_SELL":
        return "강력 매도"
    return "분석 불가"


def translate_market_regime(regime: str) -> str:
    normalized = str(regime or "").strip().lower()
    if normalized == "trend":
        return "추세장"
    if normalized == "range":
        return "횡보장"
    return "미확인"


def extract_company_reference_text(
    blocks: list[dict[str, Any]],
    company: dict[str, str],
    max_chars: int,
) -> str:
    keywords = {
        str(company.get("company_name_ko") or "").strip(),
        str(company.get("company_name") or "").strip(),
        str(company.get("ticker") or "").strip(),
    }
    keywords = {item for item in keywords if item}
    matches = []
    for block in blocks:
        text = str(block.get("text", "")).strip()
        if not text:
            continue
        sentences = split_text_into_sentences(text)
        sentence_matches = [
            sentence
            for sentence in sentences
            if any(keyword in sentence for keyword in keywords)
        ]
        if sentence_matches:
            matches.extend(sentence_matches)
            continue
        if any(keyword in text for keyword in keywords):
            matches.append(text)
    return compact_text(" ".join(matches), max_chars=max_chars)


def blocks_to_plain_text(blocks: list[dict[str, Any]], max_chars: int) -> str:
    lines: list[str] = []
    for block in blocks:
        text = str(block.get("text", "")).strip()
        if not text:
            continue
        if block.get("type") == "bullet":
            lines.append(f"• {text}")
        else:
            lines.append(text)
    return compact_text("\n\n".join(lines), max_chars=max_chars)


def compact_text(text: str, max_chars: int) -> str:
    normalized = " ".join(str(text or "").replace("\r", "\n").split())
    if len(normalized) <= max_chars:
        return normalized
    clipped = normalized[: max(0, max_chars - 1)].rstrip()
    return f"{clipped}…"


def usable_section_width(section: Any) -> Any:
    return section.page_width - section.left_margin - section.right_margin


def split_text_into_sentences(text: str) -> list[str]:
    cleaned = " ".join(str(text or "").split())
    if not cleaned:
        return []
    sentences = [item.strip() for item in re.split(r"(?<=[.!?])\s+", cleaned) if item.strip()]
    return sentences or [cleaned]


def insert_mirae_cover_page(
    document: Any,
    research_result: dict[str, Any],
    plots: list[dict[str, Any]],
    title: str,
    subtitle: str,
    company_order: list[dict[str, str]],
    fundamentals_map: dict[str, dict[str, Any]],
    opinions_map: dict[str, dict[str, Any]],
) -> None:
    article = research_result["article"]
    company = company_order[0] if company_order else {
        "ticker": "",
        "company_name_ko": "",
        "company_name": "",
        "market": "",
    }
    ticker = company.get("ticker", "")
    fundamentals = fundamentals_map.get(ticker, {})
    opinion = opinions_map.get(ticker, {})
    plot_map = {
        str(item.get("ticker", "")).strip(): item
        for item in plots
    }
    plot = plot_map.get(ticker, {})

    cover = document.add_table(rows=1, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    left_cell = cover.rows[0].cells[0]
    right_cell = cover.rows[0].cells[1]
    left_cell.width = Inches(2.0)
    right_cell.width = Inches(5.1)
    set_cell_background(left_cell, MIRAE_ORANGE)
    set_cell_background(right_cell, MIRAE_BLACK)
    remove_cell_borders(left_cell)
    remove_cell_borders(right_cell)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    clear_cell(left_cell)
    clear_cell(right_cell)

    left_intro = left_cell.paragraphs[0]
    add_styled_run(
        left_intro,
        "AI Equity Research",
        size=14,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_PAPER,
    )
    add_empty_paragraphs(left_cell, 2)
    left_meta = left_cell.add_paragraph()
    add_styled_run(
        left_meta,
        f"Equity Research\n{format_template_date(article.get('published_at'))}",
        size=10.5,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_PAPER,
    )
    add_empty_paragraphs(left_cell, 1)

    stats_table = left_cell.add_table(rows=1, cols=1)
    stats_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stats_cell = stats_table.cell(0, 0)
    set_cell_background(stats_cell, MIRAE_BEIGE)
    remove_cell_borders(stats_cell)
    set_cell_text(
        stats_cell,
        build_cover_left_metrics(opinion=opinion, fundamentals=fundamentals),
        size=9.5,
        font_name=DEFAULT_BODY_FONT,
    )

    add_empty_paragraphs(left_cell, 1)
    cover_plot_path = str(plot.get("daily_plot_path") or plot.get("plot_path") or "").strip()
    if cover_plot_path and Path(cover_plot_path).exists():
        left_cell.add_paragraph()
        picture_paragraph = left_cell.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(cover_plot_path, width=Inches(1.55))

    analyst = left_cell.add_paragraph()
    add_styled_run(
        analyst,
        f"[Source]\n{article.get('source') or 'N/A'}\n\n[Model]\n{research_result.get('model') or 'N/A'}",
        size=8.5,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_PAPER,
    )

    top_line = right_cell.paragraphs[0]
    add_styled_run(
        top_line,
        f"{ticker}  |  {company.get('company_name', '')}",
        size=9.5,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=WHITE,
    )
    add_divider_paragraph(right_cell, MIRAE_DIVIDER)

    title_paragraph = right_cell.add_paragraph()
    add_styled_run(
        title_paragraph,
        company.get("company_name_ko") or (title or "리서치 리포트"),
        size=24,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_ORANGE,
    )

    subtitle_paragraph = right_cell.add_paragraph()
    add_styled_run(
        subtitle_paragraph,
        subtitle or str(research_result["selection"].get("summary", "")).strip()[:40] or "테마 리서치",
        size=15,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_BEIGE,
    )

    article_meta = right_cell.add_paragraph()
    add_styled_run(
        article_meta,
        f"기사 제목: {article.get('title') or 'N/A'}",
        size=9.5,
        font_name=DEFAULT_BODY_FONT,
        color=WHITE,
    )
    add_divider_paragraph(right_cell, MIRAE_DIVIDER)

    quick_view = right_cell.add_paragraph()
    add_styled_run(
        quick_view,
        f"{len(company_order)}개 종목 Quick View",
        size=11.5,
        bold=True,
        font_name=DEFAULT_BODY_FONT,
        color=MIRAE_BEIGE,
    )

    summary_box = right_cell.add_table(rows=1, cols=1)
    summary_box.alignment = WD_TABLE_ALIGNMENT.LEFT
    summary_cell = summary_box.cell(0, 0)
    set_cell_background(summary_cell, MIRAE_LIGHT_BLACK)
    set_cell_text(
        summary_cell,
        str(research_result["selection"].get("summary", "")).strip() or "요약 정보가 없습니다.",
        size=10,
        font_name=DEFAULT_BODY_FONT,
        color=WHITE,
    )
    set_cell_border(summary_cell, bottom={"color": MIRAE_DIVIDER, "sz": 6})

    right_cell.add_paragraph()
    quick_table = right_cell.add_table(rows=1, cols=4)
    quick_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    quick_table.autofit = True
    headers = ["종목", "의견", "종가", "핵심지표"]
    for idx, label in enumerate(headers):
        cell = quick_table.rows[0].cells[idx]
        set_cell_background(cell, MIRAE_BEIGE)
        set_cell_text(cell, label, bold=True, size=9, font_name=DEFAULT_BODY_FONT)

    for item in company_order[:3]:
        row = quick_table.add_row().cells
        item_fundamentals = fundamentals_map.get(item["ticker"], {})
        item_opinion = opinions_map.get(item["ticker"], {})
        metrics = item_fundamentals.get("valuation_metrics", {})
        set_cell_background(row[0], MIRAE_PAPER)
        set_cell_background(row[1], MIRAE_PAPER)
        set_cell_background(row[2], MIRAE_PAPER)
        set_cell_background(row[3], MIRAE_PAPER)
        set_cell_text(row[0], f"{item['company_name_ko']}\n{item['ticker']}", size=8.8)
        set_cell_text(row[1], str(item_opinion.get("opinion") or "N/A"), size=8.8)
        set_cell_text(
            row[2],
            format_financial_metric("close", item_fundamentals.get("latest_close")),
            size=8.8,
        )
        key_metric = first_available_metric(metrics)
        set_cell_text(row[3], key_metric, size=8.8)

    right_cell.add_paragraph()


def populate_word_report_template(
    document: Any,
    research_result: dict[str, Any],
    plots: list[dict[str, Any]],
    title: str,
    subtitle: str,
    company_order: list[dict[str, str]],
    fundamentals_map: dict[str, dict[str, Any]],
    technical_map: dict[str, dict[str, Any]],
    opinions_map: dict[str, dict[str, Any]],
) -> None:
    article = research_result["article"]
    template_date = format_template_date(article.get("published_at"))
    summary_text = str(research_result["selection"].get("summary", "")).strip()
    title_text = title or subtitle or article.get("title") or "리서치 리포트"

    replace_token_in_document(document, "{Title}", title_text)
    replace_token_in_document(document, "{date}", template_date)
    replace_token_in_document(document, "{summary_text}", summary_text)
    replace_token_in_document(document, "{Insert Opinion summary Table Here}", "")
    clear_placeholder_tokens(document)

    if len(document.tables) >= 1:
        populate_template_summary_table(document.tables[0], summary_text)
    if len(document.tables) >= 2:
        populate_template_plot_table(
            table=document.tables[1],
            company_order=company_order,
            plots=plots,
            fundamentals_map=fundamentals_map,
            technical_map=technical_map,
            opinions_map=opinions_map,
        )
    if len(document.tables) >= 3:
        populate_template_context_table(
            table=document.tables[2],
            article=article,
            selection_summary=summary_text,
            title_text=title_text,
        )


def populate_template_summary_table(table: Any, summary_text: str) -> None:
    if not table.rows or not table.rows[0].cells:
        return

    set_cell_text(
        table.rows[0].cells[0],
        f"Summary\n\n{summary_text or '요약 정보가 없습니다.'}",
        bold=False,
        size=11,
        font_name=None,
    )
    for cell in table.rows[0].cells[1:]:
        set_cell_text(cell, "", font_name=None, size=None)


def populate_template_plot_table(
    table: Any,
    company_order: list[dict[str, str]],
    plots: list[dict[str, Any]],
    fundamentals_map: dict[str, dict[str, Any]],
    technical_map: dict[str, dict[str, Any]],
    opinions_map: dict[str, dict[str, Any]],
) -> None:
    if len(table.rows) < 4:
        return

    plot_map = {
        str(item.get("ticker", "")).strip(): item
        for item in plots
    }
    display_columns = [0, 1, 2]

    for cell in table.rows[0].cells:
        set_cell_text(cell, "", font_name=None, size=None)
    set_cell_text(
        table.rows[0].cells[0],
        "Selected Company Charts",
        bold=True,
        size=12,
        font_name=None,
    )

    for column_index in range(len(table.rows[1].cells)):
        company = company_order[column_index] if column_index < len(company_order) else None
        title_cell = table.rows[1].cells[column_index]
        plot_cell = table.rows[2].cells[column_index]
        context_cell = table.rows[3].cells[column_index]

        if company is None or column_index not in display_columns:
            set_cell_text(title_cell, "", font_name=None, size=None)
            set_cell_text(plot_cell, "", font_name=None, size=None)
            set_cell_text(context_cell, "", font_name=None, size=None)
            continue

        ticker = company["ticker"]
        plot = plot_map.get(ticker, {})
        fundamentals = fundamentals_map.get(ticker, {})
        technical = technical_map.get(ticker, {})
        opinion = opinions_map.get(ticker, {})

        set_cell_text(
            title_cell,
            f"{company['company_name_ko']} ({ticker})",
            bold=True,
            size=10,
            font_name=None,
        )
        inserted_daily = insert_image_into_cell(
            plot_cell,
            str(plot.get("daily_plot_path") or plot.get("plot_path") or "").strip(),
            width=Inches(2.0),
        )
        inserted_intraday = append_image_to_cell(
            plot_cell,
            str(plot.get("intraday_plot_path") or "").strip(),
            width=Inches(2.0),
        )
        if not inserted_daily and not inserted_intraday:
            set_cell_text(plot_cell, "차트 없음", font_name=None, size=10)

        set_cell_text(
            context_cell,
            build_template_plot_context(
                fundamentals=fundamentals,
                technical=technical,
                opinion=opinion,
                plot=plot,
            ),
            font_name=None,
            size=9,
        )


def populate_template_context_table(
    table: Any,
    article: dict[str, Any],
    selection_summary: str,
    title_text: str,
) -> None:
    if len(table.rows) < 3:
        return

    set_cell_text(table.rows[0].cells[0], "Theme", bold=True, size=12, font_name=None)
    set_cell_text(table.rows[0].cells[1], "", font_name=None, size=None)
    set_cell_text(table.rows[0].cells[2], "Context", bold=True, size=12, font_name=None)

    set_cell_text(
        table.rows[1].cells[0],
        (
            f"{title_text}\n"
            f"출처: {article.get('source') or 'N/A'}\n"
            f"발행일: {article.get('published_at') or 'N/A'}\n\n"
            f"{selection_summary or '요약 정보가 없습니다.'}"
        ),
        size=9,
        font_name=None,
    )

    for row in table.rows[2:]:
        for cell in row.cells:
            set_cell_text(cell, "", font_name=None, size=None)


def build_template_plot_context(
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
    opinion: dict[str, Any],
    plot: dict[str, Any],
) -> str:
    parts = [
        build_company_exhibit_summary(
            fundamentals=fundamentals,
            technical=technical,
            opinion=opinion,
        ) or "요약 정보가 없습니다.",
    ]
    daily_comment = str(plot.get("daily_chart_comment") or "").strip()
    intraday_comment = str(plot.get("intraday_chart_comment") or "").strip()
    if daily_comment:
        parts.append(f"[중기 차트 LLM 분석]\n{daily_comment}")
    if intraday_comment:
        parts.append(f"[단기 차트 LLM 분석]\n{intraday_comment}")
    return "\n\n".join(parts)


def detach_trailing_notice_table(document: Any) -> Optional[Any]:
    if not document.tables:
        return None

    last_table = document.tables[-1]
    first_cell_text = str(last_table.cell(0, 0).text or "").strip()
    if "Notice" not in first_cell_text:
        return None

    notice_copy = deepcopy(last_table._element)
    last_table._element.getparent().remove(last_table._element)
    return notice_copy


def replace_token_in_document(document: Any, token: str, replacement: str) -> None:
    for paragraph in iter_document_paragraphs(document):
        if token not in paragraph.text:
            continue
        paragraph.text = paragraph.text.replace(token, replacement)


def clear_placeholder_tokens(document: Any) -> None:
    placeholder_tokens = (
        "{Header}",
        "{Plot_title}",
        "{Plot}",
        "{Context}",
    )
    for token in placeholder_tokens:
        replace_token_in_document(document, token, "")


def iter_document_paragraphs(document: Any) -> list[Any]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def format_template_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return datetime.now().strftime("%Y.%m.%d")

    normalized = text[:10].replace("-", ".").replace("/", ".")
    return normalized


def insert_investment_opinion_table(
    document: Any,
    opinions: list[dict[str, Any]],
    use_template_styles: bool = False,
) -> None:
    add_section_heading(document, "투자의견 요약", level=2)

    if not opinions:
        add_body_paragraph(document, "투자의견 요약을 생성할 종목이 없습니다.")
        return

    table = document.add_table(rows=1, cols=len(OPINION_TABLE_COLUMNS))
    if use_template_styles:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = table.rows[0].cells
    for index, (_, label) in enumerate(OPINION_TABLE_COLUMNS):
        set_cell_background(header_cells[index], MIRAE_BLACK if not use_template_styles else MIRAE_BEIGE)
        set_cell_text(
            header_cells[index],
            label,
            bold=True,
            size=9.5,
            font_name=None if use_template_styles else DEFAULT_BODY_FONT,
            color=WHITE if not use_template_styles else None,
        )

    for row_index, company in enumerate(opinions):
        row_cells = table.add_row().cells
        for index, (key, _) in enumerate(OPINION_TABLE_COLUMNS):
            value = company.get(key)
            if key in {"score", "confidence"}:
                text = format_scalar(value, decimals=2)
            else:
                text = str(value or "N/A")
            if not use_template_styles:
                set_cell_background(
                    row_cells[index],
                    MIRAE_PAPER if row_index % 2 == 0 else WHITE,
                )
            set_cell_text(
                row_cells[index],
                text,
                size=9,
                font_name=None if use_template_styles else DEFAULT_BODY_FONT,
            )

    document.add_paragraph()


def insert_company_exhibits(
    document: Any,
    company_order: list[dict[str, str]],
    plot_map: dict[str, dict[str, Any]],
    fundamentals_map: dict[str, dict[str, Any]],
    technical_map: dict[str, dict[str, Any]],
    opinions_map: dict[str, dict[str, Any]],
    use_template_styles: bool = False,
) -> None:
    for company in company_order:
        ticker = company["ticker"]
        plot = plot_map.get(ticker, {})
        fundamentals = fundamentals_map.get(ticker, {})
        technical = technical_map.get(ticker, {})
        opinion = opinions_map.get(ticker, {})

        subheading = (
            f"{company['company_name_ko']} ({company['company_name']}) "
            f"[{company['market']}] {ticker}"
        ).strip()
        add_section_heading(document, subheading, level=3, use_template_styles=use_template_styles)

        summary = build_company_exhibit_summary(
            fundamentals=fundamentals,
            technical=technical,
            opinion=opinion,
        )
        if summary:
            add_body_paragraph(document, summary, use_template_styles=use_template_styles)

        daily_plot_path = str(plot.get("daily_plot_path") or plot.get("plot_path") or "").strip()
        if daily_plot_path and Path(daily_plot_path).exists():
            add_section_heading(document, "3개월 일봉 주가 차트", level=4, use_template_styles=use_template_styles)
            document.add_picture(daily_plot_path, width=Inches(6.3))
            add_body_paragraph(document, build_plot_caption(plot, chart_kind="daily"), use_template_styles=use_template_styles)
            add_chart_comment_section(
                document,
                "중기 차트 LLM 분석",
                plot.get("daily_chart_comment"),
                use_template_styles=use_template_styles,
            )

        intraday_plot_path = str(plot.get("intraday_plot_path") or "").strip()
        if intraday_plot_path and Path(intraday_plot_path).exists():
            add_section_heading(document, "1분봉 주가/거래량 차트", level=4, use_template_styles=use_template_styles)
            document.add_picture(intraday_plot_path, width=Inches(6.3))
            add_body_paragraph(document, build_plot_caption(plot, chart_kind="intraday"), use_template_styles=use_template_styles)
            add_chart_comment_section(
                document,
                "단기 차트 LLM 분석",
                plot.get("intraday_chart_comment"),
                use_template_styles=use_template_styles,
            )

        if not (
            daily_plot_path
            and Path(daily_plot_path).exists()
            or intraday_plot_path
            and Path(intraday_plot_path).exists()
        ):
            add_body_paragraph(
                document,
                "주가 차트를 생성할 수 있는 데이터가 충분하지 않습니다.",
                use_template_styles=use_template_styles,
            )

        add_section_heading(document, "재무 분석 표", level=4, use_template_styles=use_template_styles)
        insert_financial_analysis_table(document, fundamentals)
        document.add_paragraph()


def insert_financial_analysis_table(
    document: Any,
    fundamentals: dict[str, Any],
) -> None:
    metrics = fundamentals.get("valuation_metrics", {})
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = table.rows[0].cells
    for cell, label in zip(header_cells, ("지표", "값", "지표", "값")):
        set_cell_background(cell, MIRAE_BLACK)
        set_cell_text(cell, label, bold=True, color=WHITE)

    pairs = list(FINANCIAL_TABLE_ORDER)
    for index in range(0, len(pairs), 2):
        row_cells = table.add_row().cells
        left_key, left_label = pairs[index]
        shaded = MIRAE_PAPER if (index // 2) % 2 == 0 else WHITE
        for cell in row_cells:
            set_cell_background(cell, shaded)
        set_cell_text(row_cells[0], left_label)
        set_cell_text(row_cells[1], format_financial_metric(left_key, metrics.get(left_key)))

        if index + 1 < len(pairs):
            right_key, right_label = pairs[index + 1]
            set_cell_text(row_cells[2], right_label)
            set_cell_text(row_cells[3], format_financial_metric(right_key, metrics.get(right_key)))
        else:
            set_cell_text(row_cells[2], "")
            set_cell_text(row_cells[3], "")

    source_paragraph = document.add_paragraph()
    add_styled_run(
        source_paragraph,
        f"재무 기준일: {fundamentals.get('statement_date') or 'N/A'} | 출처: {fundamentals.get('source') or 'OpenDART'}",
        size=9,
        italic=True,
        color=MIRAE_SOFT_GRAY,
    )


def add_chart_comment_section(
    document: Any,
    title: str,
    comment: Any,
    use_template_styles: bool = False,
) -> None:
    text = str(comment or "").strip()
    if not text:
        return
    add_section_heading(document, title, level=5, use_template_styles=use_template_styles)
    add_body_paragraph(document, text, use_template_styles=use_template_styles)


def build_company_exhibit_summary(
    fundamentals: dict[str, Any],
    technical: dict[str, Any],
    opinion: dict[str, Any],
) -> str:
    fragments: list[str] = []

    opinion_label = str(opinion.get("opinion", "")).strip()
    if opinion_label:
        fragments.append(f"현재 시스템 의견은 {opinion_label}입니다.")

    score = opinion.get("score")
    if score is not None:
        fragments.append(f"점수는 {format_scalar(score, decimals=2)}입니다.")

    regime = str(technical.get("signal_context", {}).get("market_regime", "")).strip()
    if regime:
        fragments.append(f"기술적 국면은 {regime}로 분류됩니다.")

    latest_close = fundamentals.get("latest_close")
    if latest_close is not None:
        fragments.append(f"최신 종가는 {format_financial_metric('close', latest_close)}입니다.")

    return " ".join(fragments)


def parse_markdown_report(markdown: str) -> tuple[str, str, list[dict[str, Any]]]:
    title = ""
    subtitle = ""
    sections: list[dict[str, Any]] = []
    current_section: Optional[dict[str, Any]] = None
    paragraph_buffer: list[str] = []
    section_started = False

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer, current_section
        if current_section is None or not paragraph_buffer:
            paragraph_buffer = []
            return
        current_section["blocks"].append(
            {"type": "paragraph", "text": " ".join(paragraph_buffer).strip()}
        )
        paragraph_buffer = []

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue

        if line.startswith("# "):
            flush_paragraph()
            title = line[2:].strip()
            continue

        if line.startswith("## ") and not subtitle and title and not section_started:
            flush_paragraph()
            subtitle = line[3:].strip()
            continue

        if line.startswith("## "):
            flush_paragraph()
            current_section = {
                "heading": line[3:].strip(),
                "blocks": [],
            }
            sections.append(current_section)
            section_started = True
            continue

        if line.startswith("### "):
            flush_paragraph()
            if current_section is not None:
                current_section["blocks"].append(
                    {"type": "subheading", "text": line[4:].strip()}
                )
            continue

        if line.startswith("- "):
            flush_paragraph()
            if current_section is not None:
                current_section["blocks"].append(
                    {"type": "bullet", "text": line[2:].strip()}
                )
            continue

        paragraph_buffer.append(line)

    flush_paragraph()
    return title, subtitle, sections


def add_markdown_blocks(document: Any, blocks: list[dict[str, Any]]) -> None:
    for block in blocks:
        block_type = block.get("type")
        text = str(block.get("text", "")).strip()
        if not text:
            continue
        if block_type == "paragraph":
            add_body_paragraph(document, text, use_template_styles=template_styles_available(document))
        elif block_type == "bullet":
            style_name = "List Bullet" if has_style(document, "List Bullet") else None
            paragraph = document.add_paragraph(style=style_name)
            add_styled_run(
                paragraph,
                text,
                size=None if template_styles_available(document) else 11,
                font_name=None if template_styles_available(document) else DEFAULT_BODY_FONT,
                color=None if template_styles_available(document) else MIRAE_LIGHT_BLACK,
            )
        elif block_type == "subheading":
            add_section_heading(document, text, level=3, use_template_styles=template_styles_available(document))


def add_section_heading(
    document: Any,
    text: str,
    level: int = 1,
    use_template_styles: bool = False,
) -> None:
    style_name = None
    if use_template_styles:
        style_name = {
            1: "Heading 1",
            2: "Heading 2",
            3: "Heading 3",
        }.get(min(level, 3), "Heading 3")
        if not has_style(document, style_name):
            style_name = None

    paragraph = document.add_paragraph(style=style_name)
    if level == 1:
        size = 15
    elif level == 2:
        size = 13
    elif level == 3:
        size = 12
    else:
        size = 11
    add_styled_run(
        paragraph,
        text,
        size=None if use_template_styles else size,
        bold=not use_template_styles or level <= 3,
        font_name=None if use_template_styles else DEFAULT_BODY_FONT,
        color=None if use_template_styles else MIRAE_ORANGE,
    )
    if not use_template_styles:
        add_divider(document)


def add_body_paragraph(document: Any, text: str, use_template_styles: bool = False) -> None:
    style_name = "Normal" if use_template_styles and has_style(document, "Normal") else None
    paragraph = document.add_paragraph(style=style_name)
    paragraph.paragraph_format.space_after = Pt(7)
    add_styled_run(
        paragraph,
        text,
        size=None if use_template_styles else 11,
        font_name=None if use_template_styles else DEFAULT_BODY_FONT,
        color=None if use_template_styles else MIRAE_LIGHT_BLACK,
    )


def add_styled_run(
    paragraph: Any,
    text: str,
    *,
    size: Optional[float],
    bold: bool = False,
    italic: bool = False,
    font_name: Optional[str] = DEFAULT_BODY_FONT,
    color: Optional[str] = None,
) -> Any:
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if font_name:
        run.font.name = font_name
    if color and RGBColor is not None:
        try:
            run.font.color.rgb = RGBColor.from_string(color)
        except Exception:
            pass
    if qn is not None and font_name:
        try:
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        except Exception:
            pass
    return run


def configure_document_styles(document: Any) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.22)

    normal_style = document.styles["Normal"]
    normal_style.font.name = DEFAULT_BODY_FONT
    normal_style.font.size = Pt(11)
    if qn is not None:
        try:
            normal_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), DEFAULT_BODY_FONT)
        except Exception:
            pass


def set_cell_text(
    cell: Any,
    text: str,
    bold: bool = False,
    size: Optional[float] = 10,
    font_name: Optional[str] = DEFAULT_BODY_FONT,
    color: Optional[str] = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    add_styled_run(
        paragraph,
        text,
        size=size,
        bold=bold,
        font_name=font_name,
        color=color,
    )


def insert_image_into_cell(cell: Any, image_path: str, width: Any) -> bool:
    cell.text = ""
    if not image_path or not Path(image_path).exists():
        return False

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(image_path, width=width)
    return True


def append_image_to_cell(cell: Any, image_path: str, width: Any) -> bool:
    if not image_path or not Path(image_path).exists():
        return False

    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(image_path, width=width)
    return True


def clear_cell(cell: Any) -> None:
    cell.text = ""
    if not cell.paragraphs:
        cell.add_paragraph()


def add_empty_paragraphs(cell: Any, count: int) -> None:
    for _ in range(count):
        cell.add_paragraph()


def add_divider(document: Any) -> None:
    divider_table = document.add_table(rows=1, cols=1)
    divider_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    divider_cell = divider_table.cell(0, 0)
    set_cell_text(divider_cell, "", size=1)
    set_cell_border(divider_cell, bottom={"color": MIRAE_DIVIDER, "sz": 6})
    remove_cell_borders(divider_cell, keep_bottom=True)


def add_divider_paragraph(cell: Any, color: str) -> None:
    divider_table = cell.add_table(rows=1, cols=1)
    divider_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    divider_cell = divider_table.cell(0, 0)
    set_cell_text(divider_cell, "", size=1, font_name=None)
    set_cell_border(divider_cell, bottom={"color": color, "sz": 6})
    remove_cell_borders(divider_cell, keep_bottom=True)


def set_cell_background(cell: Any, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) if qn is not None else None
    if shd is None and OxmlElement is not None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    if shd is not None:
        shd.set(qn("w:fill"), color)


def set_cell_margins(
    cell: Any,
    top: int = 80,
    bottom: int = 80,
    start: int = 80,
    end: int = 80,
) -> None:
    if OxmlElement is None or qn is None:
        return

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for edge, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def remove_cell_borders(cell: Any, keep_bottom: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders")) if qn is not None else None
    if borders is None and OxmlElement is not None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    if borders is None or OxmlElement is None:
        return

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        if keep_bottom and edge == "bottom":
            continue
        element.set(qn("w:val"), "nil")


def set_cell_border(cell: Any, **kwargs: dict[str, str]) -> None:
    if OxmlElement is None or qn is None:
        return

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("left", "top", "right", "bottom", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "color", "space"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))
        if "val" not in edge_data:
            element.set(qn("w:val"), "single")


def format_financial_metric(key: str, value: Any) -> str:
    numeric = try_float(value)
    if numeric is None:
        return "N/A"
    numeric = round(numeric, 2)
    if key in PERCENT_FIELDS:
        return f"{numeric:.2f}%"
    if key == "shares_outstanding":
        return format_korean_share_count(numeric)
    if key in LARGE_NUMBER_FIELDS:
        return format_korean_large_number(numeric)
    if key in {"eps", "bps", "cash_dps", "per", "pbr", "pcr", "ev_to_ebitda"}:
        return f"{numeric:,.2f}"
    if key == "close":
        return f"{numeric:,.2f}"
    return f"{numeric:,.4f}"


def format_korean_large_number(value: float) -> str:
    sign = "-" if value < 0 else ""
    absolute = abs(value)
    trillion = int(absolute // 1_000_000_000_000)
    remainder = absolute - trillion * 1_000_000_000_000
    hundred_million = remainder / 100_000_000

    if trillion > 0:
        if hundred_million >= 1:
            return f"{sign}{trillion:,}조 {hundred_million:,.0f}억"
        return f"{sign}{trillion:,}조"
    if absolute >= 100_000_000:
        return f"{sign}{absolute / 100_000_000:,.0f}억"
    if absolute >= 10_000:
        return f"{sign}{absolute / 10_000:,.0f}만"
    return f"{sign}{absolute:,.0f}"


def format_korean_share_count(value: float) -> str:
    sign = "-" if value < 0 else ""
    absolute = abs(value)
    if absolute >= 100_000_000:
        return f"{sign}{absolute / 100_000_000:,.2f}억주"
    if absolute >= 10_000:
        return f"{sign}{absolute / 10_000:,.0f}만주"
    return f"{sign}{absolute:,.0f}주"


def format_scalar(value: Any, decimals: int = 2) -> str:
    numeric = try_float(value)
    if numeric is None:
        return "N/A"
    return f"{numeric:.{decimals}f}"


def build_cover_left_metrics(
    opinion: dict[str, Any],
    fundamentals: dict[str, Any],
) -> str:
    metrics = fundamentals.get("valuation_metrics", {})
    lines = [
        f"투자의견        {str(opinion.get('opinion') or 'N/A')}",
        f"종가            {format_financial_metric('close', fundamentals.get('latest_close'))}",
        f"PER             {format_financial_metric('per', metrics.get('per'))}",
        f"PBR             {format_financial_metric('pbr', metrics.get('pbr'))}",
        f"ROE             {format_financial_metric('roe', metrics.get('roe'))}",
        f"시가총액        {format_financial_metric('market_cap', metrics.get('market_cap'))}",
    ]
    return "\n".join(lines)


def first_available_metric(metrics: dict[str, Any]) -> str:
    for key, label in (
        ("per", "PER"),
        ("pbr", "PBR"),
        ("roe", "ROE"),
        ("operating_margin", "OPM"),
        ("market_cap", "시가총액"),
    ):
        value = metrics.get(key)
        formatted = format_financial_metric(key, value)
        if formatted != "N/A":
            return f"{label} {formatted}"
    return "N/A"


def try_float(value: Any) -> Optional[float]:
    try:
        if value is None:
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def build_document_path(output_dir: Path, article: dict[str, Any]) -> Path:
    title = str(article.get("title", "")).strip() or "research_report"
    slug = "".join(char if char.isalnum() else "_" for char in title).strip("_")
    slug = slug[:80] or "research_report"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"{timestamp}_{slug}.docx"


def resolve_word_report_template_path(template_path: Optional[str]) -> Optional[Path]:
    if not template_path:
        return None

    candidate = Path(template_path).expanduser()
    if not candidate.is_absolute():
        candidate = (Path.cwd() / candidate).resolve()
    else:
        candidate = candidate.resolve()

    if candidate.exists():
        return candidate
    raise WordReportError(f"Word template not found: {candidate}")


def has_style(document: Any, style_name: str) -> bool:
    try:
        document.styles[style_name]
        return True
    except Exception:
        return False


def template_styles_available(document: Any) -> bool:
    return bool(getattr(document, "_cam_template_mode", False))


def ensure_python_docx_available() -> None:
    if Document is None:
        raise WordReportError(
            "The 'python-docx' package is not installed. Run `pip install -r requirements.txt`."
        )
